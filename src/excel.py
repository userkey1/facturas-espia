# -*- coding: utf-8 -*-
"""
Funciones para guardar los datos capturados en un Excel.
"""

import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .config import EXCEL_PATH, COLUMNAS_EXCEL


def guardar_en_excel(datos: dict, archivo=EXCEL_PATH):
    """
    Agrega una fila al Excel. Lo crea si no existe.
    Mantiene el orden de columnas definido en config.COLUMNAS_EXCEL.
    """
    archivo = str(archivo)
    os.makedirs(os.path.dirname(archivo), exist_ok=True)

    # Respetamos el orden de columnas
    fila = {col: datos.get(col, "") for col in COLUMNAS_EXCEL}
    df_nuevo = pd.DataFrame([fila])

    # Si existe el archivo, lo abrimos y agregamos
    if os.path.exists(archivo):
        try:
            df_prev = pd.read_excel(archivo)
            df = pd.concat([df_prev, df_nuevo], ignore_index=True)
        except Exception:
            # Si el archivo está corrupto, lo recreamos
            df = df_nuevo
    else:
        df = df_nuevo

    df.to_excel(archivo, index=False)

    # Le damos formato al encabezado
    try:
        wb = load_workbook(archivo)
        ws = wb.active
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Ancho automático
        for col in ws.columns:
            max_len = max(len(str(c.value)) if c.value else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 3

        wb.save(archivo)
    except Exception as e:
        print(f"⚠️ No se pudo formatear el Excel: {e}")